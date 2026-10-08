"""The IRP research step: irp_sheet.py reads a saved copy of the US IRPs tab and the
export, build_state_brief.py --irp turns it into tasks and the statewide block, the
assembler marks records that come from a plan, and the workbook gains the irp_box and
irp_notes_draft sheets. Everything runs on synthetic data under tmp_path: no sheet read,
no network, nothing near the GEM database."""
import csv
import json
import subprocess
import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parent.parent
for p in (ROOT, ROOT / "scripts"):
    if str(p) not in sys.path:
        sys.path.insert(0, str(p))

import irp_sheet  # noqa: E402
import build_review_package as brp  # noqa: E402
from review_app import review_data  # noqa: E402

PLAN = "https://www.georgiapower.com/content/dam/irp/2025-irp-main-document.pdf"
DOCKET = "https://psc.ga.gov/search/facts-document/?documentId=12345"
EIA = "https://www.eia.gov/electricity/data/eia860m/xls/august_generator2026.xlsx"

TAB = [
    ["State", "Utility", "IRP year", "File", "Notes", "Updated IRP"],
    ["Alabama"],
    ["", "Alabama Power", "2025", "https://www.alabamapower.com/irp-2025.pdf",
     "Q2 2026 nothing new. Q4 2025 read the 2025 plan.", ""],
    ["Georgia", "", "", "", "", "", "https://psc.ga.gov/"],
    ["", "Georgia Power", "2025 Draft IRP", PLAN + "\n" + DOCKET + " (see chapter 7)",
     "Q2 2026 Draft IRP filed in January; proposes 1,500 MW of new gas CTs at Plant Yates. "
     "Next IRP due 2028.\nQ4 2025 Nothing new.", "yes"],
    ["", "Oglethorpe Power", "", "", "Q1 2026 no public IRP.", ""],
    ["Indiana", "", "", "", "", ""],
    ["", "Duke Energy Indiana", "2024", "https://www.duke-energy.com/irp-2024.pdf", "Q4 2025 read.", ""],
]

HEADER = ["Last Updated", "Researcher", "Research status", "Wiki URL", "Country/Area", "State/Province",
          "Plant name", "Other Name(s)", "Unit name", "Fuel", "Fuel Data Source", "Number Of Engines",
          "Capacity Per Engine", "Capacity (MW)", "Capacity Data Source", "Status", "Status Detail",
          "Status Data Source", "IRP", "Latest Activity", "Latest Activity Data Source", "Start year",
          "Start Year Data Source", "Retired year", "Retired Year Data Source", "Planned retire",
          "Planned Retire Data Source", "Cancellation year", "Cancellation year Data Source",
          "Turbine/Engine Technology", "Turbine/Engine Technology Data Source",
          "Equipment Manufacturer/Model", "Turbine/Engine Equipment Data Source", "CHP", "CHP Data Source",
          "Owner(s)", "Owners Data Source", "Operator(s)", "Operators Data Source", "Parent",
          "Latitude", "Longitude", "Location accuracy", "Location Data Source", "City",
          "GEM location ID", "GEM unit ID", "Captive industry use", "Captive industry type",
          "Captive non-industry use", "Captive Data Source", "Conversion/replacement?",
          "Conversion/replacement Data Source", "Other IDs (location)", "Other IDs (unit)"]


def unit(pid, uid, plant, name, status, mw, owner, irp="no", start=""):
    r = {h: "" for h in HEADER}
    r.update({"Country/Area": "United States", "State/Province": "Georgia", "Plant name": plant,
              "Unit name": name, "Fuel": "fossil gas: natural gas", "Capacity (MW)": mw, "Status": status,
              "IRP": irp, "Start year": start, "Owner(s)": owner, "Operator(s)": owner,
              "Parent": "Southern Company", "GEM location ID": pid, "GEM unit ID": uid,
              "Latitude": "33.4", "Longitude": "-84.9", "Location accuracy": "exact",
              "Status Data Source": EIA, "Capacity Data Source": EIA})
    return r


def write_csv(path):
    rows = [
        unit("L1", "G1", "Plant Yates power station", "CT1", "announced", "500", "Georgia Power Company [100%]"),
        unit("L1", "G2", "Plant Yates power station", "6", "operating", "350", "Georgia Power Company [100%]", start="1974"),
        unit("L2", "G3", "Plant McIntosh power station", "10", "operating", "600", "Georgia Power Company [100%]", start="2005"),
        unit("L3", "G4", "Georgia Power IRP Project 1 power station", "1", "pre-construction", "800",
             "Georgia Power Company [100%]", irp="yes"),
        unit("L4", "G5", "Smarr Energy Facility", "1", "operating", "250", "Oglethorpe Power Corporation [100%]", start="1999"),
    ]
    with open(path, "w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=HEADER)
        w.writeheader()
        w.writerows(rows)
    Path(path).with_suffix(".colmap.json").write_text("{}")   # the header is re-derived from the csv
    return path


def test_parse_tab_forward_fills_states_and_splits_notes():
    states = irp_sheet.parse_tab(TAB)
    assert set(states) == {"alabama", "georgia", "indiana"}
    ga = states["georgia"]
    assert ga["links"] == ["https://psc.ga.gov/"]          # the regulator link on the state row
    assert [r["utility"] for r in ga["rows"]] == ["Georgia Power", "Oglethorpe Power"]
    gp = ga["rows"][0]
    assert gp["links"] == [PLAN, DOCKET] and "chapter 7" in gp["file_note"]
    assert gp["is_draft"] is True
    assert [n["quarter"] for n in gp["notes"]] == ["Q2 2026", "Q4 2025"]
    assert gp["latest_note"]["quarter"] == "Q2 2026" and "Plant Yates" in gp["latest_note"]["text"]
    assert gp["next_irp_due"] == "2028"
    assert gp["mw_mentions"] == ["1,500 MW"] or "1,500" in " ".join(gp["mw_mentions"])
    assert ga["rows"][1]["is_draft"] is False and ga["rows"][1]["links"] == []


def test_parse_tab_refuses_a_changed_header():
    with pytest.raises(SystemExit):
        irp_sheet.parse_tab([["State", "Company", "Year"], ["Georgia"]])


def test_build_state_matches_plants_and_narrows_tasks(tmp_path):
    csv_path = write_csv(tmp_path / "export.csv")
    data = irp_sheet.build_state("Georgia", irp_sheet.parse_tab(TAB), csv_path, "2026-10-07")
    assert data["skipped"] is False and data["export_has_irp_column"] is True
    assert data["where"] == {"kind": "state", "name": "Georgia", "slug": "georgia",
                             "country": "United States", "postal": "ga"}
    gp = data["utilities"][0]
    names = {p["plant_name"]: p for p in gp["plants"]}
    assert set(names) == {"Plant Yates power station", "Plant McIntosh power station",
                          "Georgia Power IRP Project 1 power station"}
    # the placeholder plant is matched by name too and carries the ticked box
    ph = names["Georgia Power IRP Project 1 power station"]
    assert "plant name" in ph["matched_by"] and ph["irp_box_units"] == 1 and ph["own_task"]
    # an announced unit makes the plant task-worthy; an operating-only fleet does not
    assert names["Plant Yates power station"]["own_task"] is True
    assert names["Plant McIntosh power station"]["own_task"] is False
    assert {t["plant_id"] for t in data["tasks"]} == {"L1", "L3"}
    t = next(t for t in data["tasks"] if t["plant_id"] == "L1")
    assert t["unit_id"] is None and t["kind"] == "fix" and t["checks"] == [37] and t["source"] == "irp"
    assert 'listed on the US IRPs tab as "2025 Draft IRP"' in t["task"] and "draft draft" not in t["task"].lower()
    assert PLAN in t["task"]
    # the units already ticked, with the utility they belong to
    assert [u["unit_id"] for u in data["irp_units"]] == ["G4"]
    assert data["irp_units"][0]["matched_utility"] == "Georgia Power"
    # Oglethorpe: matched by owner, no plan link, operating fleet: no task, listed in the brief
    og = data["utilities"][1]
    assert [p["plant_name"] for p in og["plants"]] == ["Smarr Energy Facility"]
    assert "Smarr Energy Facility" in data["brief"] and "compare it here" in data["brief"]
    assert "Integrated resource plans" in data["brief"] or "resource plan" in data["brief"].lower()
    # the brief is plain language: no em-dashes or arrows
    assert "—" not in data["brief"] and "->" not in data["brief"]


def test_build_state_skips_a_state_with_no_rows(tmp_path):
    csv_path = write_csv(tmp_path / "export.csv")
    data = irp_sheet.build_state("Maryland", irp_sheet.parse_tab(TAB), csv_path, "2026-10-07")
    assert data["skipped"] is True and data["utilities"] == [] and data["tasks"] == []
    assert "skipped" in data["brief"].lower() or "no row" in data["brief"].lower()


def test_main_from_json_writes_work_files(tmp_path):
    csv_path = write_csv(tmp_path / "export.csv")
    tab = tmp_path / "irp_tab.json"
    tab.write_text(json.dumps(TAB))
    out = tmp_path / "work"
    irp_sheet.main(["--state", "Georgia", "--state", "Indiana", "--csv", str(csv_path),
                    "--from-json", str(tab), "--out-dir", str(out)])
    assert (out / "irp_us-georgia.json").exists() and (out / "irp_us-georgia.md").exists()
    ind = json.loads((out / "irp_us-indiana.json").read_text())
    assert ind["skipped"] is False and ind["utilities"][0]["plants"] == []   # no Indiana plant in the csv


def run(script, *args):
    return subprocess.run([sys.executable, str(ROOT / "scripts" / script), *args],
                          capture_output=True, text=True, cwd=ROOT / "scripts")


def test_irp_step_end_to_end(tmp_path):
    """brief builder --irp -> _irp.md and index.irp; a shard and _state.json whose sources are the
    plan -> staged records flagged irp, staged_newplants.json, irp_summary.json; the workbook's
    irp sheets and the review page data."""
    import openpyxl
    csv_path = write_csv(tmp_path / "export.csv")
    tab = tmp_path / "irp_tab.json"
    tab.write_text(json.dumps(TAB))
    work = tmp_path / "work"
    irp_sheet.main(["--state", "Georgia", "--csv", str(csv_path), "--from-json", str(tab), "--out-dir", str(work)])
    irp_json = work / "irp_us-georgia.json"
    batch = tmp_path / "batches" / "us-ga"
    r = run("build_state_brief.py", "--state", "Georgia", "--mode", "update", "--csv", str(csv_path),
            "--batch", str(batch), "--irp", str(irp_json), "--worklist", str(tmp_path / "no-worklist.csv"))
    assert r.returncode == 0, r.stdout + r.stderr
    assert "IRP step: 2 utility plan(s), 2 plant task(s), 1 unit(s) already carry the IRP box" in r.stdout
    index = json.loads((batch / "briefs" / "_index.json").read_text())
    assert index["irp"]["links"] == [PLAN, DOCKET] and index["irp"]["irp_box_units"] == ["G4"]
    assert (batch / "briefs" / "_irp.md").exists()
    yates = next(p for p in index["plants"] if p["plant_id"] == "L1")
    assert any(t["checks"] == [37] for t in yates["plant_tasks"])
    assert "integrated resource plan" in (batch / "briefs" / "L1.md").read_text().lower()

    # the research: a shard for Plant Yates with a finding from the plan, and the statewide file
    ver = [{"url": PLAN, "ok": True, "contains_value": True, "name_found": True, "note": ""}]
    (batch / "shards" / "L1.json").write_text(json.dumps({
        "meta": {"plant_id": "L1", "plant_name": "Plant Yates power station", "state": "Georgia",
                 "mode": "update", "model": "test", "generated": "2026-10-07", "done": True,
                 "urls_attempted": 1, "urls_verified": 1},
        "units": [{"gem_unit_id": "G1", "unit_name": "CT1",
                   "findings": {"Capacity (MW)": {"value": "750", "refs": [PLAN], "verifications": ver,
                                                  "tier": "high", "independent": False,
                                                  "note": "The draft plan sizes the new CTs at 750 MW."}},
                   "not_found": [], "notes": ""}],
        "plant_findings": {}, "qa": [], "monitor": [], "newunits": [], "entities": [], "source_log": []}))
    (batch / "shards" / "_state.json").write_text(json.dumps({
        "meta": {"irp_summary": {"Georgia Power": "The draft plan adds 1,500 MW of gas turbines at Plant Yates.",
                                 "Oglethorpe Power": "No public plan."}},
        "qa": [], "monitor": [],
        "newplants": [
            {"plant_name": "Plant Yates CT expansion", "fields": {"Owner(s)": "Georgia Power Company [100%]"},
             "refs": {"Owners Data Source": [PLAN]}, "verifications": ver, "tier": "medium",
             "note": "Proposed in the draft plan; no certificate filed yet.",
             "units": [{"unit_name": "CT2", "fields": {"Status": "proposed", "Capacity (MW)": "750"},
                        "refs": {"Status Data Source": [PLAN]}, "verifications": ver}]},
            {"plant_name": "Oglethorpe peaker idea", "note": "One line in a board agenda, nothing more."}],
        "newunits": [{"gem_plant_id": "L2", "unit_name": "11", "note": "A second combined-cycle block is hinted at."}],
    }))
    r = run("assemble_state.py", "--batch", str(batch), "--csv", str(csv_path), "--stamp", "20261007_1200_ET")
    assert r.returncode == 0, r.stdout + r.stderr
    staging = batch / "staging"
    lanes = brp.load_lanes(staging)
    up = lanes["updates"]
    assert len(up) == 1 and up[0]["irp"] is True and up[0]["fields"] == {"Capacity (MW)": "750"}
    assert (staging / "staged_newplants.json").exists()
    np_ = lanes["newplants"]
    assert len(np_) == 1 and np_[0]["irp"] is True and np_[0]["record_id"] == "us-ga:new:plant-yates-ct-expansion"
    assert np_[0]["units"][0]["unit_name"] == "CT2"
    # the thin new-plant lead and the new-unit lead are watch items, not edits
    mon = lanes["monitor"]
    kinds = {m["plant_name"]: m["monitor_kind"] for m in mon}
    assert kinds == {"Oglethorpe peaker idea": "new_to_tracker", "Plant McIntosh power station": "existing_plant"}
    assert all(m["monitor_reason"].startswith("New ") for m in mon)
    assert "newunits" not in lanes or lanes["newunits"] == []
    summary = json.loads((staging / "irp_summary.json").read_text())
    assert summary["utilities"]["Georgia Power"].startswith("The draft plan") and summary["skipped"] is False

    # the workbook: irp_box pairs each plan-sourced record with the unit's box; irp_notes_draft per utility
    irp = brp.load_irp(irp_json, staging)
    assert brp.is_us(lanes)
    out = tmp_path / "out.xlsx"
    brp.build_xlsx(lanes, out, export_csv=csv_path, irp=irp)
    wb = openpyxl.load_workbook(out)
    assert "irp_box" in wb.sheetnames and "irp_notes_draft" in wb.sheetnames

    def rows(ws):
        hdr = [c.value for c in ws[1]]
        return [dict(zip(hdr, r)) for r in ws.iter_rows(min_row=2, values_only=True)]
    box = rows(wb["irp_box"])
    by_unit = {(r["plant"], r["gem_unit_id"]): r for r in box}
    yates_row = by_unit[("Plant Yates power station", "G1")]
    assert yates_row["irp_box_now"] == "not ticked" and yates_row["action"].startswith("tick the IRP box")
    assert yates_row["utility"] == "Georgia Power" and yates_row["checklist_rows"] and "37" in yates_row["checklist_rows"]
    new_row = next(r for r in box if r["plant"] == "Plant Yates CT expansion")
    assert new_row["irp_box_now"] == "new row" and new_row["utility"] == "Georgia Power"
    ref = by_unit[("Georgia Power IRP Project 1 power station", "G4")]
    assert ref["irp_box_now"] == "already ticked" and ref["what_it_is"] == "unit already in the database"
    notes = {r["utility"]: r for r in rows(wb["irp_notes_draft"])}
    assert set(notes) == {"Georgia Power", "Oglethorpe Power"}
    gp = notes["Georgia Power"]
    assert gp["is_draft"] == "yes" and gp["next_irp_due"] == "2028" and gp["records_from_this_plan"] == 2
    assert gp["draft_notes_line"].startswith("Q") and "Staged in this batch" in gp["draft_notes_line"]
    assert "Plant Yates power station (Capacity (MW))" in gp["draft_notes_line"]
    assert "new row for Plant Yates CT expansion" in gp["draft_notes_line"]
    assert notes["Oglethorpe Power"]["records_from_this_plan"] == 0
    assert "already tracked" in notes["Oglethorpe Power"]["draft_notes_line"]
    ev = tmp_path / "out.md"
    brp.build_evidence(lanes, ev, "20261007_1200_ET", "us-ga", "update", export_csv=csv_path, irp=irp)
    text = ev.read_text()
    assert "## Utility resource plans (checklist row 37)" in text and "from the Georgia Power plan" in text

    # the review page data: lines and items carry irp, units carry the box state
    data = review_data.build([staging], export_csv=csv_path)
    yates = next(p for p in data["plants"] if p["pid"] == "L1")
    assert yates["lines"][0]["irp"] is True
    assert {u["gem_unit_id"]: u["irp_box"] for u in yates["units"]} == {"G1": "no"}
    newp = next(p for p in data["plants"] if p["name"] == "Plant Yates CT expansion")
    assert newp["lines"][0]["irp"] is True and newp["lines"][0]["kind"] == "new_row"
    watch = next(p for p in data["plants"] if p["pid"] == "L2")
    assert watch["items"][0]["irp"] is False


def test_build_refuses_irp_outside_the_us(tmp_path):
    lanes = {"updates": [{"record_id": "x", "country": "Germany", "fields": {}}]}
    assert brp.is_us(lanes) is False
    with pytest.raises(SystemExit):
        brp.load_irp(tmp_path / "missing.json")
