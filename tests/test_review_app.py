"""Tests for the GOGPT review app (review_app/) and the build script's --decisions consumer.

Everything runs on a throwaway staging dir under tmp_path: no real batch, no network, and
nothing that could touch the GEM database (there is no code path for that anyway)."""
import json
import os
import sys
import threading
import urllib.request
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parent.parent
for p in (ROOT, ROOT / "scripts"):
    if str(p) not in sys.path:
        sys.path.insert(0, str(p))

from review_app import review_data, server, store  # noqa: E402
import build_review_package as brp  # noqa: E402


def envelope(lane, records):
    return {"meta": {"lane": lane, "scope": {"country": "United States", "state": "Maryland", "quarter": "q4-2026"},
                     "generated": "2026-10-02T12:00:00-04:00", "counts": {"records": len(records)}},
            "records": records}


EIA = "https://www.eia.gov/electricity/data/eia860m/xls/august_generator2026.xlsx"
PR = "https://www.example.com/press/2026-05-18-brandywine"


def write_staging(d):
    """A small batch: a fill, a change on Status with one publisher, a plant-wide re-verified
    value, one concern, one monitor item with no GEM plant id."""
    d.mkdir(parents=True)
    updates = [
        {"record_id": "L1:G1:start-year", "gem_plant_id": "L1", "plant_name": "Brandywine power facility",
         "gem_unit_id": "G1", "unit_name": "F701", "country": "United States",
         "fields": {"Start year": "1996"}, "refs": {"Start Year Data Source": [EIA]},
         "verifications": [{"url": EIA, "ok": True, "contains_value": True, "name_found": True, "note": ""}],
         "current": {"Start year": ""}, "tier": "high", "verdict": "fill",
         "researcher_notes": "The EIA table for August 2026 lists 1996.", "action": "Type 1996 in Start year."},
        {"record_id": "L1:G1:status", "gem_plant_id": "L1", "plant_name": "Brandywine power facility",
         "gem_unit_id": "G1", "unit_name": "F701", "country": "United States",
         "fields": {"Status": "mothballed"}, "refs": {"Status Data Source": [PR]},
         "verifications": [{"url": PR, "ok": True, "contains_value": True, "name_found": True, "note": ""}],
         "current": {"Status": "operating"}, "tier": "medium", "verdict": "change"},
        {"record_id": "L1:G1:latitude", "gem_plant_id": "L1", "plant_name": "Brandywine power facility",
         "gem_unit_id": "G1", "unit_name": "F701", "country": "United States",
         "fields": {"Latitude": "38.70"}, "refs": {"Location Data Source": [EIA]},
         "verifications": [{"url": EIA, "ok": True, "contains_value": True, "name_found": True, "note": ""}],
         "current": {"Latitude": "38.70"}, "tier": "medium", "verdict": "match", "reverified": True,
         "applies_to_all_units": True, "sibling_unit_ids": ["G2"], "sibling_current": {"G2": "38.70"}},
    ]
    qa = [{"record_id": "L1:G1:qa-unverified-value-1", "gem_plant_id": "L1", "plant_name": "Brandywine power facility",
           "gem_unit_id": "G1", "unit_name": "F701", "country": "United States", "fields": {},
           "concern_type": "unverified_value", "recommendation": "Check the capacity against the EIA table.",
           "researcher_notes": "The capacity on the press release differs from the export."}]
    monitor = [{"record_id": "us-md:plant:monitor-1", "gem_plant_id": "", "plant_name": "Chesapeake repowering (Acme)",
                "country": "United States", "fields": {}, "item": "new plant candidate", "recheck_by": "2027-01",
                "refs": {"links": [PR]}, "researcher_notes": "Announced, no permit filed."}]
    newplants = [{"record_id": "us-md-eia999:plant:newplant", "gem_plant_id": "", "plant_name": "Severn barge station",
                  "country": "United States", "fields": {"Owner(s)": "Acme Power LLC [100%]", "City": "Baltimore"},
                  "refs": {"Owners Data Source": [PR]},
                  "verifications": [{"url": PR, "ok": True, "contains_value": True, "name_found": True, "note": ""}],
                  "tier": "medium", "researcher_notes": "Listed by EIA and the grid operator, missing from GEM.",
                  "action": "Create the plant, then add the two barge units below.",
                  "units": [
                      {"unit_name": "Barge 1", "fields": {"Status": "operating", "Start year": "1971"},
                       "refs": {"Status Data Source": [EIA]},
                       "verifications": [{"url": EIA, "ok": True, "contains_value": True, "name_found": True, "note": ""}],
                       "action": "Add unit Barge 1."},
                      {"unit_name": "Barge 2", "fields": {"Status": "retired", "Retired year": "2022"},
                       "refs": {"Status Data Source": [EIA]},
                       "verifications": [{"url": EIA, "ok": True, "contains_value": True, "name_found": True, "note": ""}]},
                  ]}]
    (d / "staged_updates.json").write_text(json.dumps(envelope("updates", updates)))
    (d / "staged_qa.json").write_text(json.dumps(envelope("qa", qa)))
    (d / "staged_monitor.json").write_text(json.dumps(envelope("monitor", monitor)))
    (d / "staged_newplants.json").write_text(json.dumps(envelope("newplants", newplants)))
    return d


@pytest.fixture
def staging(tmp_path):
    return write_staging(tmp_path / "us-md" / "staging")


@pytest.fixture
def data(staging, tmp_path):
    csv = tmp_path / "export.csv"
    csv.write_text("GEM unit ID,GEM location ID,Start year,Start Year Data Source,Status,Status Data Source,Latitude,Location Data Source\n"
                   "G1,L1,,https://old.example.com/a,operating,,38.70,\n")
    return review_data.build([staging], export_csv=csv)


# ---- initials -------------------------------------------------------------------------

@pytest.mark.parametrize("who,want", [("Baird Langenbrunner", "BL"), ("baird.langenbrunner@globalenergymonitor.org", "BL"),
                                      ("amalia", "A"), ("BL", "BL"), ("", "")])
def test_initials(who, want):
    assert store.initials(who) == want


# ---- dataset ---------------------------------------------------------------------------

def test_dataset_shape(data, staging):
    assert data["scope"]["states"] == ["Maryland"] and data["scope"]["quarter"] == "q4-2026"
    assert data["dirs"] == [str(staging)]          # outside the repo: an absolute label
    plants = {p["pid"]: p for p in data["plants"]}
    # the candidate card drops the "(Acme)" parenthetical so the same candidate staged twice shares a card
    assert set(plants) == {"L1", "new:chesapeake-repowering", "new:severn-barge-station"}
    assert data["us"] is True
    # the nine checklist groups, plus "other" because the fixture's concern names no column or kind
    assert [g["id"] for g in data["groups"]] == [1, 2, 3, 4, 5, 6, 7, 8, 9, "other"]
    for p in data["plants"]:
        for o in p["lines"] + p["items"]:
            assert o["group"] and o["group_label"] and isinstance(o["checks"], list)
    p = plants["L1"]
    assert [u["gem_unit_id"] for u in p["units"]] == ["G1"]
    lines = {l["record_id"]: l for l in p["lines"]}
    assert lines["L1:G1:start-year"]["kind"] == "fill" and lines["L1:G1:start-year"]["severity"] == "major"
    assert lines["L1:G1:start-year"]["default"] == "accept"
    assert lines["L1:G1:start-year"]["current_ref"] == ["https://old.example.com/a"]   # from the export csv
    assert lines["L1:G1:start-year"]["proposed_refs"] == [EIA]
    assert lines["L1:G1:start-year"]["verifications"][EIA]["contains_value"] is True
    st = lines["L1:G1:status"]
    assert st["kind"] == "change" and st["publishers"] == 1 and st["default"] == "hold"
    pw = lines["L1:G1:latitude"]
    assert pw["kind"] == "plant" and pw["reverified"] and pw["severity"] == "minor"
    assert pw["sibling_unit_ids"] == ["G2"]
    items = {it["record_id"]: it for it in p["items"]}
    assert items["L1:G1:qa-unverified-value-1"]["kind"] == "concern"
    assert items["L1:G1:qa-unverified-value-1"]["concern_type"] == "unverified_value"
    cand = plants["new:chesapeake-repowering"]
    assert cand["items"][0]["kind"] == "monitor" and cand["items"][0]["links"] == [PR]
    # a watch item with no GEM id is new to the tracker and lands in the new-plants group
    assert cand["items"][0]["monitor_kind"] == "new_to_tracker" and cand["items"][0]["group"] == 8
    assert all(l["key"] == f"{staging}::{l['record_id']}" for l in p["lines"])


def test_new_plant_line_carries_its_units(data):
    """A newplants record is one decision for the plant and its nested unit rows: the units
    ride along on the line, and their links count toward the line's publishers and marks."""
    npl = next(p for p in data["plants"] if p["pid"] == "new:severn-barge-station")
    assert len(npl["lines"]) == 1 and not npl["items"]
    l = npl["lines"][0]
    assert l["kind"] == "new_row" and l["lane"] == "newplants" and l["plant_name"] == "Severn barge station"
    assert l["proposed_values"] == {"Owner(s)": "Acme Power LLC [100%]", "City": "Baltimore"}
    assert [u["unit_name"] for u in l["units"]] == ["Barge 1", "Barge 2"]
    assert l["units"][0]["fields"] == {"Status": "operating", "Start year": "1971"}
    assert l["units"][0]["refs_by_col"] == {"Status Data Source": [EIA]} and l["units"][0]["action"] == "Add unit Barge 1."
    assert l["proposed_refs"] == [PR, EIA] and l["publishers"] == 2
    assert l["verifications"][EIA]["contains_value"] is True and PR in l["verifications"]
    assert l["default"] == "hold"


def test_severity_rule():
    assert review_data.line_kind({"verdict": "match"}) == "reverified"
    assert review_data.line_kind({"delete": True}) == "delete"
    assert review_data.line_kind({"applies_to_all_units": True}) == "plant"
    assert review_data.line_kind({"verdict": "change"}) == "change"
    assert review_data.line_kind({}) == "fill"


# ---- store -----------------------------------------------------------------------------

def test_store_round_trip(data, staging):
    key = f"{staging}::L1:G1:start-year"
    dirs = store.dir_paths(data)
    saved = store.decide([{"key": key, "decision": "accept"}], data, "BL", dirs)
    assert len(saved) == 1
    r = saved[0]
    assert {"key", "dir", "pid", "record_id", "column", "kind", "decision", "suggested_value", "note",
            "reviewer", "ts", "undecided"} <= set(r)
    assert r["decision"] == "accept" and r["reviewer"] == "BL" and r["undecided"] is False
    log = staging / "review_log.jsonl"
    assert log.exists() and len(log.read_text().splitlines()) == 1
    derived = json.loads((staging / "review_decisions.json").read_text())
    assert derived["decisions"][key]["decision"] == "accept"
    # suggest needs a value or a note
    with pytest.raises(store.Invalid):
        store.decide([{"key": key, "decision": "suggest"}], data, "BL", dirs)
    store.decide([{"key": key, "decision": "suggest", "suggested_value": "1997", "note": "typo"}], data, "BL", dirs)
    store.overlay(data, dirs)
    line = next(l for p in data["plants"] for l in p["lines"] if l["key"] == key)
    assert line["decision"] == "suggest" and line["suggested_value"] == "1997" and line["reviewed"] is True
    # undo: the log grows, the overlay clears
    store.decide([{"key": key, "undo": True}], data, "BL", dirs)
    assert len(log.read_text().splitlines()) == 3
    store.overlay(data, dirs)
    assert line["decision"] is None and line["reviewed"] is False
    # unknown key, bad decision
    with pytest.raises(store.Invalid):
        store.decide([{"key": "nope", "decision": "accept"}], data, "BL", dirs)
    with pytest.raises(store.Invalid):
        store.decide([{"key": key, "decision": "maybe"}], data, "BL", dirs)


def test_store_items(data, staging):
    key = f"{staging}::L1:G1:qa-unverified-value-1"
    dirs = store.dir_paths(data)
    saved = store.record_items([{"key": key, "call": "confirmed", "note": "stands"}], data, "BL", dirs)
    assert saved[0]["call"] == "confirmed" and saved[0]["kind"] == "concern"
    with pytest.raises(store.Invalid):
        store.record_items([{"key": key, "call": "noted"}], data, "BL", dirs)   # not in the concern vocabulary
    store.overlay(data, dirs)
    it = next(i for p in data["plants"] for i in p["items"] if i["key"] == key)
    assert it["call"] == "confirmed" and it["call_note"] == "stands" and it["decided_by"] == "BL"


def test_store_watch_calls(data, staging):
    """A watch item takes the four watchlist calls, never accept/reject; removing needs a note."""
    key = f"{staging}::us-md:plant:monitor-1"
    dirs = store.dir_paths(data)
    assert store.ITEM_CALLS["monitor"] == ("add_to_database", "hold", "possible_updates", "remove")
    with pytest.raises(store.Invalid):
        store.record_items([{"key": key, "call": "confirmed"}], data, "BL", dirs)
    with pytest.raises(store.Invalid, match="needs a note"):
        store.record_items([{"key": key, "call": "remove"}], data, "BL", dirs)
    with pytest.raises(store.Invalid, match="web address"):
        store.record_items([{"key": key, "call": "hold", "reference": "not a url"}], data, "BL", dirs)
    saved = store.record_items([{"key": key, "call": "add_to_database", "note": "permit is final",
                                 "reference": "https://example.com/permit"}], data, "BL", dirs)
    assert saved[0]["call"] == "add_to_database" and saved[0]["reference"] == "https://example.com/permit"
    store.overlay(data, dirs)
    it = next(i for p in data["plants"] for i in p["items"] if i["key"] == key)
    assert it["call"] == "add_to_database" and it["call_reference"] == "https://example.com/permit"


def test_store_flags(data, staging):
    """Ask-the-PM flags sit beside decisions: their own key space, on/off, latest wins."""
    dirs = store.dir_paths(data)
    key = f"{staging}::L1:G1:start-year"
    store.decide([{"key": key, "decision": "accept"}], data, "BL", dirs)
    saved = store.record_flags([{"key": key, "on": True, "note": "is 1996 right?"}], data, "BL", dirs)
    assert saved[0]["key"] == "pm::" + key and saved[0]["flag"] == "pm" and saved[0]["on"] is True
    assert saved[0]["record_id"] == "L1:G1:start-year" and saved[0]["kind"] == "fill"   # the line's own kind
    # a flag on the whole plant names the plant, not a staged record
    pf = store.record_flags([{"pid": "L1", "dir": str(staging), "on": True, "note": "whole plant looks odd"}], data, "BL", dirs)
    assert pf[0]["key"] == store.plant_flag_key(str(staging), "L1") and pf[0]["record_id"] is None and pf[0]["kind"] == "plant"
    with pytest.raises(store.Invalid):
        store.record_flags([{"key": "nope", "on": True}], data, "BL", dirs)
    store.overlay(data, dirs)
    p = next(p for p in data["plants"] if p["pid"] == "L1")
    line = next(l for l in p["lines"] if l["key"] == key)
    assert line["decision"] == "accept" and line["pm_flag"] is True and line["pm_note"] == "is 1996 right?" and line["pm_by"] == "BL"
    assert p["pm_flag"] is True and p["pm_note"] == "whole plant looks odd"
    other = next(l for l in p["lines"] if l["key"] != key)
    assert other["pm_flag"] is False
    # off again: the log grows, the overlay clears, the decision stays
    store.record_flags([{"key": key, "on": False}], data, "BL", dirs)
    store.overlay(data, dirs)
    assert line["pm_flag"] is False and line["pm_note"] == "" and line["decision"] == "accept"
    # the build script keeps decisions and flags apart
    decisions = brp.load_decisions(staging)
    assert set(decisions) == {"L1:G1:start-year"} and decisions["L1:G1:start-year"]["decision"] == "accept"
    flags = brp.load_flags(staging)
    assert set(flags) == {"L1"} and flags["L1"]["note"] == "whole plant looks odd"


# ---- server ----------------------------------------------------------------------------

def test_server_round_trip(data, staging, tmp_path):
    path = tmp_path / "review_data.json"
    path.write_text(json.dumps(data))
    app = server.App(path, "Baird Langenbrunner")
    srv = server.make_server(app, port=0)
    port = srv.server_address[1]
    t = threading.Thread(target=srv.serve_forever, daemon=True)
    t.start()
    base = f"http://127.0.0.1:{port}"

    def call(path, body=None):
        req = urllib.request.Request(base + path, data=json.dumps(body).encode() if body is not None else None,
                                     headers={"Content-Type": "application/json"} if body is not None else {})
        try:
            with urllib.request.urlopen(req) as r:
                return r.status, json.loads(r.read().decode())
        except urllib.error.HTTPError as e:
            return e.code, json.loads(e.read().decode())
    try:
        st, who = call("/api/whoami")
        assert st == 200 and who == {"reviewer": "BL", "caps": {"decide": True}}
        key = f"{staging}::L1:G1:status"
        st, body = call("/api/decide", [{"key": key, "decision": "hold"}])
        assert st == 200 and body["saved"][0]["decision"] == "hold" and body["saved"][0]["reviewer"] == "BL"
        st, body = call("/api/decide", [{"key": "nope", "decision": "hold"}])
        assert st == 400 and "unknown key" in body["error"]
        st, body = call("/api/item", [{"key": f"{staging}::L1:G1:qa-unverified-value-1", "call": "needs_research"}])
        assert st == 200 and body["saved"][0]["call"] == "needs_research"
        st, body = call("/api/flag", [{"key": key, "on": True, "note": "ask about the source"}])
        assert st == 200 and body["saved"][0]["flag"] == "pm" and body["saved"][0]["on"] is True
        st, d = call("/api/data")
        line = next(l for p in d["plants"] for l in p["lines"] if l["key"] == key)
        assert line["decision"] == "hold" and line["decided_by"] == "BL"
        assert line["pm_flag"] is True and line["pm_note"] == "ask about the source"
        with urllib.request.urlopen(base + "/") as r:
            assert r.status == 200 and b"GOGPT reviewer" in r.read()
    finally:
        srv.shutdown()
        srv.server_close()
    assert (staging / "review_log.jsonl").exists()


def test_server_refuses_non_loopback():
    with pytest.raises(ValueError):
        server.ensure_loopback("0.0.0.0")


# ---- build script: --decisions -----------------------------------------------------------

def test_build_apply_decisions(data, staging):
    dirs = store.dir_paths(data)
    k = lambda rid: f"{staging}::{rid}"
    store.decide([{"key": k("L1:G1:start-year"), "decision": "accept"},
                  {"key": k("L1:G1:status"), "decision": "reject", "note": "one source"},
                  {"key": k("us-md-eia999:plant:newplant"), "decision": "accept"}], data, "BL", dirs)
    store.decide([{"key": k("L1:G1:latitude"), "decision": "suggest", "suggested_value": "38.71"}], data, "BL", dirs)
    store.record_items([{"key": k("L1:G1:qa-unverified-value-1"), "call": "confirmed"}], data, "BL", dirs)
    lanes = brp.load_lanes(staging)
    assert brp.validate(lanes) == []
    decisions = brp.load_decisions(staging)
    kept, left_out, counts = brp.apply_decisions(lanes, decisions)
    assert counts == {"accept": 2, "hold": 0, "reject": 1, "suggest": 1, "undecided": 0}
    assert [r["record_id"] for r in kept["updates"]] == ["L1:G1:start-year"]
    # one call on a new plant carries its nested unit rows into the build
    assert [r["record_id"] for r in kept["newplants"]] == ["us-md-eia999:plant:newplant"]
    assert [u["unit_name"] for u in kept["newplants"][0]["units"]] == ["Barge 1", "Barge 2"]
    assert sorted(r["record_id"] for _, r, _ in left_out) == ["L1:G1:latitude", "L1:G1:status"]
    assert kept["qa"][0]["review_call"]["call"] == "confirmed"
    assert len(kept["monitor"]) == 1 and "review_call" not in kept["monitor"][0]
    texts = {r["record_id"]: brp.describe_decision(d) for _, r, d in left_out}
    assert texts["L1:G1:status"].startswith("rejected by BL on 20") and texts["L1:G1:status"].endswith(". one source")
    assert "suggested by BL" in texts["L1:G1:latitude"] and texts["L1:G1:latitude"].endswith(": 38.71")
    assert brp.describe_decision(None) == "not decided yet in the review app"
    # the checklist carries only the accepted edit
    rows = brp.checklist_rows(kept)
    assert len(rows) == 1 and rows[0][5] == "Start year" and rows[0][7] == "1996"


def test_build_no_log(staging):
    assert brp.load_decisions(staging) is None


def test_build_watch_calls_and_pm_questions(tmp_path):
    """Watch items split by kind; the four watchlist calls feed their own sheets; ask-the-PM
    flags make a questions sheet; removed items leave the workbook; the evidence file is
    organized by checklist group."""
    import openpyxl
    staging = write_staging(tmp_path / "us-md" / "staging")
    mon = json.loads((staging / "staged_monitor.json").read_text())
    mon["records"] += [
        {"record_id": "L1:plant:monitor-2", "gem_plant_id": "L1", "plant_name": "Brandywine power facility",
         "country": "United States", "fields": {}, "monitor_reason": "The owner says a fourth unit is under study.",
         "recheck_by": "2027-03", "refs": {"links": [PR]}, "researcher_notes": "No filing yet."},
        {"record_id": "L1:plant:monitor-3", "gem_plant_id": "L1", "plant_name": "Brandywine power facility",
         "country": "United States", "fields": {}, "monitor_reason": "A battery project shares the site.",
         "recheck_by": "2027-03", "refs": {"links": [PR]}},
        {"record_id": "us-md:plant:monitor-4", "gem_plant_id": "", "plant_name": "Patapsco peaker (Beta)",
         "country": "United States", "fields": {}, "monitor_reason": "Named in a county board agenda only.",
         "recheck_by": "2027-06", "refs": {"links": [PR]}},
    ]
    mon["meta"]["counts"]["records"] = len(mon["records"])
    (staging / "staged_monitor.json").write_text(json.dumps(mon))
    csv = tmp_path / "export.csv"
    csv.write_text("GEM unit ID,GEM location ID,Start year,Start Year Data Source,Status,Status Data Source,Latitude,Location Data Source\n"
                   "G1,L1,,https://old.example.com/a,operating,,38.70,\n")
    (tmp_path / "export.colmap.json").write_text("{}")   # the header is re-derived from the csv
    data = review_data.build([staging], export_csv=csv)
    dirs = store.dir_paths(data)
    k = lambda rid: f"{staging}::{rid}"
    store.decide([{"key": k("L1:G1:start-year"), "decision": "accept"}], data, "BL", dirs)
    store.record_items([
        {"key": k("us-md:plant:monitor-1"), "call": "add_to_database", "note": "the permit is final",
         "reference": "https://example.com/permit"},
        {"key": k("L1:plant:monitor-2"), "call": "possible_updates", "note": "for the next cycle"},
        {"key": k("L1:plant:monitor-3"), "call": "remove", "note": "a battery, not a gas unit"},
        {"key": k("L1:G1:qa-unverified-value-1"), "call": "confirmed"},
    ], data, "BL", dirs)
    store.record_flags([{"key": k("L1:G1:start-year"), "on": True, "note": "is 1996 the first unit or the plant?"},
                        {"pid": "new:" + review_data.candidate_key("Chesapeake repowering (Acme)"), "dir": str(staging),
                         "on": True, "note": "do we track repowerings as new plants?"}], data, "BL", dirs)

    lanes = brp.load_lanes(staging)
    decisions = brp.load_decisions(staging)
    kept, left_out, counts = brp.apply_decisions(lanes, decisions)
    # the removed watch item leaves the kept lanes and is listed as left out
    assert {r["record_id"] for r in kept["monitor"]} == {"us-md:plant:monitor-1", "L1:plant:monitor-2", "us-md:plant:monitor-4"}
    assert any(r["record_id"] == "L1:plant:monitor-3" for _, r, _ in left_out)
    gone = next(d for _, r, d in left_out if r["record_id"] == "L1:plant:monitor-3")
    assert brp.describe_decision(gone).startswith("remove from watchlist by BL on 20") and "a battery" in brp.describe_decision(gone)
    assert brp.describe_call({"call": "add_to_database", "reviewer": "BL", "at": "2026-10-07T10:00:00-04:00",
                              "reference": "https://example.com/permit"}).startswith("incorporate into database by BL")
    assert brp.describe_call({"call": "add_to_database", "reviewer": "BL", "at": "2026-10-07T10:00:00-04:00",
                              "reference": "https://example.com/permit"}).endswith("Source given: https://example.com/permit")
    # the checklist carries the group and sheet rows beside each edit (groups are tagged at build time)
    assert brp.tag_lanes(kept, csv) is True
    rows = brp.checklist_rows(kept)
    assert len(brp.CHECKLIST_HEADER) == 14 and brp.CHECKLIST_HEADER[-3:] == ["checklist_group", "checklist_rows", "done"]
    assert len(rows) == 1 and rows[0][5] == "Start year" and rows[0][11].startswith("2.")
    flags = brp.load_flags(staging)
    assert len(flags) == 2
    qrows = brp.pm_question_rows(lanes, flags)
    assert len(qrows) == 2
    assert {q[8] for q in qrows} == {"is 1996 the first unit or the plant?", "do we track repowerings as new plants?"}
    assert "Chesapeake" in next(q[1] for q in qrows if "repowerings" in q[8])

    out = tmp_path / "out.xlsx"
    brp.build_xlsx(kept, out, export_csv=csv, flags=flags, all_lanes=lanes)
    wb = openpyxl.load_workbook(out)
    names = wb.sheetnames
    for want in ("checklist_summary", "edit_checklist", "watch_new_to_tracker", "watch_existing_plants",
                 "promote_to_database", "possible_updates_rows", "questions_for_pm"):
        assert want in names, want
    assert "monitor_list" not in names

    def col(ws, name):
        hdr = [c.value for c in ws[1]]
        return [r[hdr.index(name)] for r in ws.iter_rows(min_row=2, values_only=True)]
    assert sorted(col(wb["watch_new_to_tracker"], "plant")) == ["Chesapeake repowering (Acme)", "Patapsco peaker (Beta)"]
    assert col(wb["watch_existing_plants"], "plant") == ["Brandywine power facility"]
    assert col(wb["watch_existing_plants"], "reviewer_call")[0].startswith("send to possible updates by BL")
    assert col(wb["watch_new_to_tracker"], "reviewer_call") and any(
        (c or "").startswith("incorporate into database by BL") for c in col(wb["watch_new_to_tracker"], "reviewer_call"))
    assert col(wb["promote_to_database"], "plant") == ["Chesapeake repowering (Acme)"]
    assert col(wb["possible_updates_rows"], "plant") == ["Brandywine power facility"]
    assert col(wb["possible_updates_rows"], "reviewer_note") == ["for the next cycle"]
    assert len(list(wb["questions_for_pm"].iter_rows(min_row=2))) == 2
    summary = wb["checklist_summary"]
    assert [c.value for c in summary[1]][:3] == ["checklist_group", "row", "checkbox"]
    assert summary.max_row > 40

    ev = tmp_path / "out.md"
    brp.build_evidence(kept, ev, "20261007_1000_ET", "us-md", "update", left_out=left_out, flags=flags,
                       all_lanes=lanes, export_csv=csv)
    text = ev.read_text()
    assert "## Questions for the PM" in text and "## Checklist group" in text
    assert text.index("## Questions for the PM") < text.index("## Checklist group")
    assert "a battery, not a gas unit" in text   # the removed item and its reason
    assert "Left out by the review" in text

    # build_state_brief.py --promote: the "incorporate into database" calls become sweep work
    import build_state_brief as bsb
    existing, new = bsb.load_promotions(staging)
    assert existing == {} and len(new) == 1   # the promoted item has no GEM plant yet
    assert new[0]["plant_name"] == "Chesapeake repowering (Acme)" and new[0]["checks"] == [51]
    assert "https://example.com/permit" in new[0]["links"] and new[0]["note"] == "the permit is final"
    store.record_items([{"key": k("L1:plant:monitor-2"), "call": "add_to_database", "note": "study is funded"}], data, "BL", dirs)
    existing, new = bsb.load_promotions(staging)
    assert set(existing) == {"L1"} and len(existing["L1"]) == 1
    text, fields, checks = existing["L1"][0]
    assert text.startswith("The review on 20") and "fourth unit is under study" in text and "study is funded" in text
    assert checks == [51] and fields == list(bsb.RESEARCH_FIELDS)


# ---- single-file page: build_static.py, static_store.js, import_log.py ---------------------

from review_app import build_static, import_log  # noqa: E402
import shutil  # noqa: E402
import subprocess  # noqa: E402

NODE = shutil.which("node")

NODE_DRIVER = r"""
const fs = require("fs"), vm = require("vm");
const [, , cfgPath, jsPath] = process.argv;
const mem = {};
global.window = global;
global.localStorage = {getItem: k => (k in mem ? mem[k] : null), setItem: (k, v) => { mem[k] = v; },
                       key: i => Object.keys(mem)[i], get length() { return Object.keys(mem).length; }};
global.Intl = Intl;
window.REVIEW_STATIC = JSON.parse(fs.readFileSync(cfgPath, "utf8"));
if (process.env.EARLIER_LOG) mem["review-log:earlier-build:" + window.REVIEW_STATIC.data.dirs.join(",")] = process.env.EARLIER_LOG;
vm.runInThisContext(fs.readFileSync(jsPath, "utf8"));
const S = window.StaticStore;
(async () => {
  const d0 = await S.load();
  const line = d0.plants[0].lines.find(l => l.record_id === "L1:G1:start-year");
  const status = d0.plants[0].lines.find(l => l.record_id === "L1:G1:status");
  const item = d0.plants[0].items.find(i => i.record_id === "L1:G1:qa-unverified-value-1");
  const out = {who: await S.whoami(), caps: S.caps, before: line.decision};
  await S.decide([{key: line.key, decision: "accept"}]);
  await S.decide([{key: status.key, decision: "suggest", suggested_value: "shelved", note: "one source"}]);
  await S.decide([{key: status.key, undo: true}]);
  await S.item([{key: item.key, call: "confirmed"}]);
  try { await S.decide([{key: item.key, decision: "accept"}]); out.itemAsLine = "allowed"; } catch (e) { out.itemAsLine = e.message; }
  try { await S.decide([{key: line.key, decision: "maybe"}]); out.badDecision = "allowed"; } catch (e) { out.badDecision = e.message; }
  // watch items: the four watchlist calls, a note needed to remove; flags beside decisions
  const watch = d0.plants.flatMap(p => p.items).find(i => i.record_id === "us-md:plant:monitor-1");
  try { await S.item([{key: watch.key, call: "confirmed"}]); out.watchBad = "allowed"; } catch (e) { out.watchBad = e.message; }
  try { await S.item([{key: watch.key, call: "remove"}]); out.watchNoNote = "allowed"; } catch (e) { out.watchNoNote = e.message; }
  await S.item([{key: watch.key, call: "add_to_database", note: "permit final", reference: "https://example.com/permit"}]);
  await S.flag([{key: line.key, on: true, note: "is 1996 right?"}]);
  await S.flag([{pid: "L1", dir: line.dir, on: true, note: "whole plant"}]);
  try { await S.flag([{key: "nope", on: true}]); out.flagBad = "allowed"; } catch (e) { out.flagBad = e.message; }
  const d1 = await S.load();   // a reload lays the browser log back over the data
  const l1 = d1.plants[0].lines.find(l => l.record_id === "L1:G1:start-year");
  const s1 = d1.plants[0].lines.find(l => l.record_id === "L1:G1:status");
  const i1 = d1.plants[0].items.find(i => i.record_id === "L1:G1:qa-unverified-value-1");
  const w1 = d1.plants.flatMap(p => p.items).find(i => i.record_id === "us-md:plant:monitor-1");
  out.after = {line: [l1.decision, l1.reviewed, l1.decided_by], status: [s1.decision, s1.reviewed], item: [i1.call, i1.reviewed],
               watch: [w1.call, w1.call_reference], flag: [l1.pm_flag, l1.pm_note, d1.plants[0].pm_flag, d1.plants[0].pm_note, s1.pm_flag]};
  out.log = S.log();
  out.stored = Object.keys(mem).length;
  process.stdout.write(JSON.stringify(out));
})();
"""


def test_static_page_renders(data):
    html = build_static.render(data, store.initials("Amalia Llano"))
    assert "<style>" in html and 'href="style.css"' not in html and 'src="app.js"' not in html
    assert "window.REVIEW_STATIC = " in html and "window.StaticStore" in html
    cfg = json.loads(html.split("window.REVIEW_STATIC = ", 1)[1].split(";</script>", 1)[0].replace("<\\/", "</"))
    assert cfg["reviewer"] == "AL" and len(cfg["data"]["plants"]) == len(data["plants"])
    assert 'id="export-btn"' in html


@pytest.mark.skipif(not NODE, reason="node not installed")
def test_static_store_round_trip(data, staging, tmp_path):
    cfg = tmp_path / "cfg.json"
    cfg.write_text(json.dumps({"reviewer": "AL", "data": data}))
    driver = tmp_path / "drive.js"
    driver.write_text(NODE_DRIVER)
    r = subprocess.run([NODE, str(driver), str(cfg), str(ROOT / "review_app" / "web" / "static_store.js")],
                       capture_output=True, text=True, check=True)
    out = json.loads(r.stdout)
    assert out["who"] == "AL" and out["caps"] == {"decide": True} and out["before"] is None
    assert out["itemAsLine"].endswith("is an item, not a line")
    assert "not one of accept, hold, reject, suggest" in out["badDecision"]
    assert out["after"] == {"line": ["accept", True, "AL"], "status": [None, False], "item": ["confirmed", True],
                            "watch": ["add_to_database", "https://example.com/permit"],
                            "flag": [True, "is 1996 right?", True, "whole plant", False]}
    assert "not one of add to database".replace("add to database", "add_to_database") in out["watchBad"]
    assert "needs a note" in out["watchNoNote"] and "unknown key" in out["flagBad"]
    assert out["stored"] == 1
    log = out["log"]
    assert [r.get("decision") or r.get("call") or r.get("flag") for r in log] == ["accept", "suggest", "hold", "confirmed", "add_to_database", "pm", "pm"]
    assert log[2]["undecided"] is True and log[2]["decision"] == "hold"      # an undo keeps the line's default
    assert log[5]["key"] == "pm::" + log[0]["key"] and log[6]["key"].endswith("::plant:L1") and log[6]["record_id"] is None
    assert all(r["reviewer"] == "AL" and r["ts"][:4] == "2026" and r["dir"] == str(staging) for r in log)
    # the download goes through import_log.py into the staging dir's log, exactly like the server
    f = tmp_path / "review_log_us-md_AL.jsonl"
    f.write_text("".join(json.dumps(r) + "\n" for r in log))
    import_log.main([str(f)])
    assert len(store.read_log(staging)) == 7 and (staging / store.DERIVED_NAME).exists()
    decisions = brp.load_decisions(staging)
    assert decisions["L1:G1:start-year"]["decision"] == "accept" and decisions["L1:G1:status"]["undecided"] is True
    assert decisions["us-md:plant:monitor-1"]["call"] == "add_to_database"
    assert set(brp.load_flags(staging)) == {"L1:G1:start-year", "L1"}
    # importing the same download again appends nothing
    planned = import_log.plan(log)
    assert [len(v[1]) for v in planned.values()] == [0] and [len(v[2]) for v in planned.values()] == [7]
    # a bad watch call or a malformed flag stops the import
    with pytest.raises(SystemExit, match="not one of"):
        import_log.plan([dict(log[4], call="confirmed", ts="2026-10-07T10:00:00-04:00")])
    with pytest.raises(SystemExit, match="flag"):
        import_log.plan([dict(log[5], key=log[0]["key"])])
    # a record the batch no longer has stops the import before anything is written
    bad = dict(log[0], record_id="L1:G1:gone", key=f"{staging}::L1:G1:gone")
    with pytest.raises(SystemExit, match="not in the staged lane files"):
        import_log.plan([bad])


@pytest.mark.skipif(not NODE, reason="node not installed")
def test_static_store_carries_calls_over_a_rebuild(data, staging, tmp_path):
    """A rebuilt page starts from the calls this browser made on an earlier build, minus dropped records."""
    cfg = tmp_path / "cfg.json"
    cfg.write_text(json.dumps({"reviewer": "AL", "data": data}))
    driver = tmp_path / "drive.js"
    driver.write_text(NODE_DRIVER)
    rec = {"dir": str(staging), "decision": "reject", "reviewer": "AL", "ts": "2026-10-03T10:00:00-04:00", "undecided": False}
    earlier = [dict(rec, key=f"{staging}::L1:G1:start-year", record_id="L1:G1:start-year"),
               dict(rec, key=f"{staging}::L1:G1:gone", record_id="L1:G1:gone")]
    r = subprocess.run([NODE, str(driver), str(cfg), str(ROOT / "review_app" / "web" / "static_store.js")],
                       capture_output=True, text=True, check=True, env=dict(os.environ, EARLIER_LOG=json.dumps(earlier)))
    out = json.loads(r.stdout)
    assert out["before"] == "reject"                                   # the earlier call shows on the rebuilt page
    assert [x["record_id"] for x in out["log"]][0] == "L1:G1:start-year"
    assert all(x["record_id"] != "L1:G1:gone" for x in out["log"])     # a dropped record's call is left behind
