"""Coordinate hyperlinks, location evidence lines and the location validity check
in build_review_package."""
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
for p in (ROOT, ROOT / "scripts"):
    if str(p) not in sys.path:
        sys.path.insert(0, str(p))

import build_review_package as brp  # noqa: E402
import geo  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

PIN = "https://www.google.com/maps/search/?api=1&query=38.69512,-76.84071"
DOC = "https://example.gov/permit.pdf"


def loc_rec(**extra):
    r = {"record_id": "L1:plant:location", "gem_unit_id": "G1", "gem_plant_id": "L1",
         "plant_name": "Brandywine", "unit_name": "F701", "country": "United States",
         "verdict": "change",
         "fields": {"Latitude": "38.69512", "Longitude": "-76.84071",
                    "Location accuracy": "exact"},
         "current": {"Latitude": "38.7", "Longitude": "-76.8",
                     "Location accuracy": "approximate"},
         "refs": {"Location Data Source": [DOC, PIN]},
         "verifications": [{"url": DOC, "ok": True}, {"url": PIN, "ok": True}],
         "tier": "high", "checks": [25], "action": "Move the point.",
         "location_method": "map", "location_moved_m": 1240,
         "location_control_points": [{"feature": "stack", "px": [412, 233],
                                      "lat": 38.69, "lon": -76.84}],
         "location_fit_m": 45, "location_image": "staging/maps/L1.jpg",
         "location_needs_review": True}
    r.update(extra)
    return r


def other_rec():
    return {"record_id": "L2:G2:capacity", "gem_unit_id": "G2", "gem_plant_id": "L2",
            "plant_name": "Zeta", "unit_name": "1", "country": "United States",
            "fields": {"Capacity (MW)": "50"}, "current": {"Capacity (MW)": "40"},
            "refs": {"Capacity Data Source": [DOC]},
            "verifications": [{"url": DOC, "ok": True}], "tier": "high", "action": "x"}


def test_workbook_links(tmp_path):
    lanes = {"updates": [loc_rec(), other_rec()]}
    out = tmp_path / "a.xlsx"
    brp.build_xlsx(lanes, out)
    ws = load_workbook(out)["edit_checklist"]
    rows = {r[5].value: r for r in ws.iter_rows(min_row=2)}
    new = geo.satellite_link("38.69512", "-76.84071")
    old = geo.satellite_link("38.7", "-76.8")
    for col in ("Latitude", "Longitude"):
        r = rows[col]
        assert r[6].hyperlink.target == old
        assert r[7].hyperlink.target == new
        assert r[6].style == "Hyperlink"
        assert r[0].hyperlink is None and r[8].hyperlink is None
    assert rows["Location accuracy"][7].hyperlink is None
    assert rows["Capacity (MW)"][6].hyperlink is None
    assert rows["Capacity (MW)"][7].hyperlink is None
    assert rows["Latitude"][7].value == "38.69512"


def test_no_current_link_when_gem_has_none():
    rec = loc_rec(current={})
    rows, links = brp.checklist_rows_and_links({"updates": [rec]})
    assert len(rows) == 3
    assert all(c == 7 for (_, c) in links) and len(links) == 2
    assert brp.checklist_rows({"updates": [rec]}) == rows


def test_evidence_lines(tmp_path):
    ev = tmp_path / "e.md"
    brp.build_evidence({"updates": [loc_rec()]}, ev, "20261020_1200_ET", "us-md", "update")
    text = ev.read_text()
    new = geo.links_for("38.69512", "-76.84071")
    old = geo.links_for("38.7", "-76.8")
    assert (f"- location now: 38.7, -76.8 (approximate) [satellite]({old['satellite']}) "
            f"[pin]({old['pin']})") in text
    assert (f"- location proposed: 38.69512, -76.84071 (exact) [satellite]({new['satellite']}) "
            f"[pin]({new['pin']})") in text
    assert "- moved: 1.2 km" in text
    assert "- method: georeferenced from a map in a document" in text
    assert "- map render: [staging/maps/L1.jpg](staging/maps/L1.jpg)" in text
    assert "  - control point: stack at pixel 412,233 is 38.69, -76.84" in text
    assert "- fit error: 45 m" in text
    assert "- needs a second look: yes" in text


def test_evidence_fill_and_short_move(tmp_path):
    ev = tmp_path / "e.md"
    rec = loc_rec(current={}, location_moved_m=340, location_method="address",
                  location_needs_review=False)
    brp.build_evidence({"updates": [rec, other_rec()]}, ev, "s", "us-md", "update")
    text = ev.read_text()
    assert "- location now: none in GEM" in text
    assert "- moved: 340 m" in text
    assert "- method: placed from an address" in text
    assert "second look" not in text and "map render" not in text
    assert text.count("location proposed") == 1


def test_validate_needs_all_three():
    rec = loc_rec()
    del rec["fields"]["Longitude"]
    errs = brp.validate({"updates": [rec]})
    assert any("needs Latitude, Longitude and Location accuracy together" in e
               and "not Longitude" in e for e in errs)
    rev = loc_rec()
    del rev["fields"]["Latitude"]
    del rev["fields"]["Longitude"]
    assert any("not Latitude, Longitude" in e for e in brp.validate({"updates": [rev]}))
    assert not [e for e in brp.validate({"updates": [loc_rec()]}) if "location edit" in e]


def test_sibling_current_dict_accepted():
    rec = {"record_id": "r", "gem_unit_id": "G1", "gem_plant_id": "L1",
           "fields": {"Notes": "new\n\nold"}, "current": {"Notes": "old"},
           "sibling_current": {"G2": {"Notes": "old"}}}
    assert brp.additive_errors(rec, rec, "G1") == []
    rec["sibling_current"] = {"G2": {"Notes": "different"}}
    assert brp.additive_errors(rec, rec, "G1")
