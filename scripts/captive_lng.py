#!/usr/bin/env python3
"""
Read the captive LNG sheet for one scope and turn it into the captive-power
research step (QC/Country checklist rows 6 and 44; row 39, gas-powered data
centers, for a US state).

Usage (from scripts/):
    python captive_lng.py --state Maryland
    python captive_lng.py --country Germany
    python captive_lng.py --state Texas --refresh           # re-download both workbooks
    python captive_lng.py --state Texas --from-dir ../work/captive   # offline / tests

Reads
    the captive LNG workbook ("Use this - Captive PPs_All Regions_<date>.xlsx",
    the file checklist row 6 links; CAPTIVE_FILE below) and the LNG tracker
    download it was built from (LNG_FILE). Both are xlsx files on the work
    Drive, downloaded through the gws work profile (read-only, never a public
    export URL) into work/captive/ and kept seven days; --refresh re-downloads.
    They hold a colleague's working notes: keep them under work/, never commit.
    The captive workbook has one pair of tabs per region: "<Region> -
    Qualifying (>=20MW)" (EU and UK) or "(>=50MW)" (everywhere else), and
    "<Region> - Excluded". Terminal ID is the LNG tracker's Project ID, so the
    tracker supplies the country, state, status, owner and coordinates for
    every row, and its "Captive gas power plant GEM ID" column names the GOGPT
    plant where the LNG team already linked one.
    the GOGPT-scoped export, for the scope's plants.

Writes  work/captive_<tag>.json and .md (tag = us-<state-slug> or <country-slug>)
    where        the scope (same shape as validation_report.py / irp_sheet.py)
    read         the date the workbooks were downloaded
    threshold_mw 20 for EU and UK rows, 50 elsewhere (the tab's own figure)
    terminals    one per sheet row in scope: the sheet's columns, the tracker's
                 columns, qualifying (true for a Qualifying tab row),
                 mechanical_only (the sheet qualifies it on compressor-drive
                 turbines alone, which GEM does not count as generating
                 capacity), the GEM plants matched (by the tracker's captive
                 plant ID, by name, or by distance) and the references with
                 gem.wiki links set aside as not citable
    gem_only     GEM plants in scope that look like LNG captive plants (name
                 says LNG, or captive type says liquefaction) but match no row
    data_center  for a US state: the plants GEM already marks as data-center
                 captive, for the row 39 block
    tasks        one whole-plant task per matched GEM plant, in the shape
                 build_state_brief.py --captive reads (kind fix, checks [6, 44])
    brief        the block for the scope-wide research agent (also the .md):
                 the matched terminals, the qualifying terminals GEM lacks (new
                 plant candidates), the excluded terminals (do not add), and for
                 a US state the data-center search instructions (row 39; watch
                 list only, Baird 2026-10-08)

Rules the step carries (user rulings 2026-10-08): the only LNG captive plants
at the add threshold are the sheet's Qualifying rows (50 MW, or 20 MW in the EU
and UK); everything else from the search goes to the watch list. The gas-fired
data center search stages into the watch list only, never a new plant.
"""
from __future__ import annotations

import argparse
import csv
import datetime
import json
import math
import re
import subprocess
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))
from paths import gogpt_scoped_csv, work_dir  # noqa: E402

CAPTIVE_FILE = "19YvVT918PYHq_DkdvKWtoXYJ2Xl7l7R3"   # captive LNG workbook (xlsx on Drive)
LNG_FILE = "1iOrpj-498jLR7QSkNRljUsgmfzRYNUfO"       # LNG tracker download (xlsx on Drive)
CAPTIVE_URL = f"https://docs.google.com/spreadsheets/d/{CAPTIVE_FILE}"
LNG_URL = f"https://docs.google.com/spreadsheets/d/{LNG_FILE}"
CACHE_DAYS = 7
NEAR_KM = 5.0          # a GEM plant this close to the terminal is a candidate match
CLOSE_KM = 2.0         # this close counts as the terminal's own plant

TAB_RE = re.compile(r"^(?P<region>.+?)\s*-\s*(?:Qualifying\s*\(>=\s*(?P<mw>\d+)\s*MW\)|(?P<excl>Excluded))\s*$", re.I)
URL_RE = re.compile(r"https?://[^\s,;]+")
GEM_ID_RE = re.compile(r"\b[LG]\d{12}\b")
MECH_RE = re.compile(r"mechanical[- ]drive|shaft", re.I)
NOT_SUMMED_RE = re.compile(r"n/?a\s*\(mechanical", re.I)
SHEET_COUNTRY = {"usa": "United States", "us": "United States", "uk": "United Kingdom",
                 "puerto rico (usa)": "Puerto Rico", "russia": "Russia"}
STOP = {"lng", "terminal", "terminals", "fsru", "flng", "export", "import", "project", "facility",
        "the", "of", "de", "and", "&", "power", "station", "plant", "captive", "regasification",
        "liquefaction", "t1", "t2", "t3", "phase", "ii", "iii", "i", "cogeneration", "cogen"}
H_PLANT, H_LOC, H_UNIT, H_UNAME = "Plant name", "GEM location ID", "GEM unit ID", "Unit name"
CAPTIVE_COLS = ("Captive industry use", "Captive industry type", "Captive non-industry use")
LIQ_TYPES = ("lng", "liquefaction")


def ws(v):
    return re.sub(r"\s+", " ", str(v if v is not None else "")).strip()


def norm_key(s):
    return re.sub(r"[^a-z0-9]+", "", ws(s).lower())


def name_tokens(s):
    s = re.sub(r"\([^)]*\)", " ", ws(s).lower())
    toks = [t for t in re.split(r"[^a-z0-9&]+", s) if t and t not in STOP]
    return set(toks)


def split_links(cell):
    out = []
    for u in URL_RE.findall(str(cell or "")):
        u = u.rstrip(").,;")
        if u not in out:
            out.append(u)
    return out


def is_gem_link(u):
    return "gem.wiki" in u or "globalenergymonitor.org" in u


def num(v):
    try:
        return float(str(v).replace(",", ""))
    except (TypeError, ValueError):
        return None


def km(lat1, lon1, lat2, lon2):
    if None in (lat1, lon1, lat2, lon2):
        return None
    p = math.pi / 180
    a = (0.5 - math.cos((lat2 - lat1) * p) / 2
         + math.cos(lat1 * p) * math.cos(lat2 * p) * (1 - math.cos((lon2 - lon1) * p)) / 2)
    return 12742 * math.asin(math.sqrt(max(0.0, min(1.0, a))))


# ---------------------------------------------------------------- Drive files

def cache_dir():
    d = work_dir() / "captive"
    d.mkdir(parents=True, exist_ok=True)
    return d


def download(file_id, name, out_dir):
    """gws drive files get, alt media, into out_dir (gws writes --output relative
    to its working directory, so the call runs there)."""
    from gsheets import _env
    params = {"fileId": file_id, "supportsAllDrives": True, "alt": "media"}
    cmd = ["gws", "drive", "files", "get", "--params", json.dumps(params), "--output", name]
    p = subprocess.run(cmd, env=_env(), capture_output=True, text=True, cwd=str(out_dir))
    target = out_dir / name
    if p.returncode != 0 or not target.exists() or target.stat().st_size == 0:
        tail = " | ".join(((p.stderr or p.stdout).strip().splitlines() or ["no output"])[-3:])
        if "401" in tail or "invalid_grant" in tail or "expired" in tail.lower():
            sys.exit("ERROR: the gws-gem login has expired; ask Baird to run "
                     "`gws-gem auth login` in his shell, then retry")
        sys.exit(f"ERROR: could not download {name} ({file_id}) with gws: {tail}")
    return target


def fetch_workbooks(refresh, from_dir=None):
    """Both xlsx files, from the cache when fresh. Returns (captive, tracker, read date)."""
    d = Path(from_dir) if from_dir else cache_dir()
    meta_p = d / "meta.json"
    cap, trk = d / "captive_lng.xlsx", d / "lng_tracker.xlsx"
    meta = {}
    if meta_p.exists():
        try:
            meta = json.loads(meta_p.read_text(encoding="utf-8"))
        except json.JSONDecodeError:
            meta = {}
    if from_dir:
        for p in (cap, trk):
            if not p.exists():
                sys.exit(f"ERROR: {p} not found; run without --from-dir to download it")
        return cap, trk, meta.get("read") or datetime.date.fromtimestamp(cap.stat().st_mtime).isoformat()
    fresh = False
    if cap.exists() and trk.exists() and not refresh:
        try:
            age = (datetime.date.today() - datetime.date.fromisoformat(meta["read"])).days
            fresh = age < CACHE_DAYS
        except (KeyError, ValueError):
            fresh = False
    if not fresh:
        download(CAPTIVE_FILE, cap.name, d)
        download(LNG_FILE, trk.name, d)
        meta = {"read": datetime.date.today().isoformat(),
                "captive_file": CAPTIVE_FILE, "lng_file": LNG_FILE}
        meta_p.write_text(json.dumps(meta, indent=1), encoding="utf-8")
    return cap, trk, meta["read"]


# ---------------------------------------------------------------- the workbooks

def _rows(ws_):
    out = []
    for r in ws_.iter_rows(values_only=True):
        if any(c is not None and ws(c) for c in r):
            out.append(list(r))
    return out


def _header_index(header):
    return {norm_key(h): i for i, h in enumerate(header) if h is not None and ws(h)}


def _get(row, idx, *names):
    for n in names:
        i = idx.get(norm_key(n))
        if i is not None and i < len(row):
            return row[i]
    return None


def read_captive_sheet(path):
    """Every row of every regional tab as a dict with region, qualifying,
    threshold_mw and the sheet's own columns."""
    import warnings
    import openpyxl
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows = []
    for ws_ in wb.worksheets:
        m = TAB_RE.match(ws_.title)
        if not m:
            continue
        data = _rows(ws_)
        if len(data) < 2:
            continue
        idx = _header_index(data[0])
        if "terminalid" not in idx:
            continue
        qual = m.group("excl") is None
        thr = int(m.group("mw")) if qual else None
        for r in data[1:]:
            tid = ws(_get(r, idx, "Terminal ID"))
            if not tid.startswith("T"):
                continue
            refs = split_links(_get(r, idx, "References"))
            rows.append({
                "tab": ws_.title, "region": ws(m.group("region")), "qualifying": qual,
                "threshold_mw": thr, "terminal": ws(_get(r, idx, "Terminal")), "terminal_id": tid,
                "sheet_country": ws(_get(r, idx, "Country")), "sheet_state": ws(_get(r, idx, "State")),
                "sheet_status": ws(_get(r, idx, "Status")),
                "hardware": ws(_get(r, idx, "Hardware Type")),
                "unit_mw": ws(_get(r, idx, "Individual Unit MW")),
                "aggregate_mw": ws(_get(r, idx, "Aggregate/Total MW", "MW Figure (if any)")),
                "basis": ws(_get(r, idx, "Qualifying Basis")),
                "turbine_class": ws(_get(r, idx, "Turbine/Genset Class")),
                "reason_excluded": ws(_get(r, idx, "Reason Excluded")),
                "refs": [u for u in refs if not is_gem_link(u)],
                "gem_refs": [u for u in refs if is_gem_link(u)],
            })
    wb.close()
    if not rows:
        sys.exit(f"ERROR: no regional Qualifying/Excluded tab with a Terminal ID column in {path}")
    return rows


def read_lng_tracker(path):
    """The LNG tracker's first tab, one dict per terminal (Project ID), with the
    unit statuses, owner, parent, operator, country, state, coordinates and the
    captive GEM plant IDs or names the LNG team recorded."""
    import warnings
    import openpyxl
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws_ = wb.worksheets[0]
    data = _rows(ws_)
    if len(data) < 2:
        sys.exit(f"ERROR: the LNG tracker tab {ws_.title!r} in {path} is empty")
    idx = _header_index(data[0])
    if "projectid" not in idx:
        sys.exit(f"ERROR: no Project ID column on tab {ws_.title!r} of {path}")
    terms = {}
    for r in data[1:]:
        tid = ws(_get(r, idx, "ProjectID"))
        if not tid.startswith("T"):
            continue
        t = terms.setdefault(tid, {
            "terminal_id": tid, "name": ws(_get(r, idx, "TerminalName")),
            "country": ws(_get(r, idx, "Country/Area")), "state": ws(_get(r, idx, "State/Province")),
            "statuses": [], "facility_types": [], "owner": "", "parent": "", "operator": "",
            "lat": None, "lon": None, "location": ws(_get(r, idx, "Location")),
            "captive_gem_ids": [], "captive_names": [], "researched": [], "notes": []})
        for key, col in (("statuses", "Status"), ("facility_types", "FacilityType")):
            v = ws(_get(r, idx, col))
            if v and v not in t[key]:
                t[key].append(v)
        for key, col in (("owner", "Owner"), ("parent", "Parent"), ("operator", "Operator")):
            if not t[key]:
                t[key] = ws(_get(r, idx, col))
        if t["lat"] is None:
            t["lat"], t["lon"] = num(_get(r, idx, "Latitude")), num(_get(r, idx, "Longitude"))
        cap = ws(_get(r, idx, "Captive gas power plant GEM ID (if applicable)"))
        if cap and cap not in ("#N/A", "N/A", "n/a"):
            ids = GEM_ID_RE.findall(cap)
            for i in ids:
                if i not in t["captive_gem_ids"]:
                    t["captive_gem_ids"].append(i)
            if not ids and cap not in t["captive_names"]:
                t["captive_names"].append(cap)
        rs = ws(_get(r, idx, "Researched?"))
        if rs and rs not in t["researched"]:
            t["researched"].append(rs)
        note = ws(_get(r, idx, "Additional notes"))
        if note and note not in ("#N/A",) and note not in t["notes"]:
            t["notes"].append(note)
    wb.close()
    return terms, ws_.title


# ---------------------------------------------------------------- the export

def load_export(csv_path, scope):
    plants = {}
    with open(csv_path, encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f)
        hdr = reader.fieldnames or []
        for h in (H_PLANT, H_LOC, H_UNIT, "State/Province", "Country/Area"):
            if h not in hdr:
                sys.exit(f"ERROR: column {h!r} missing from {csv_path}")
        for r in reader:
            if ws(r.get("Country/Area")).lower() != scope["country"].lower():
                continue
            if scope["kind"] == "state" and ws(r.get("State/Province")).lower() != scope["name"].lower():
                continue
            plants.setdefault(ws(r.get(H_LOC)), []).append(r)
    return plants


def plant_summary(pid, rows):
    first = rows[0]
    caps = [num(r.get("Capacity (MW)")) for r in rows]
    return {"plant_id": pid, "plant_name": ws(first.get(H_PLANT)),
            "state": ws(first.get("State/Province")), "lat": num(first.get("Latitude")),
            "lon": num(first.get("Longitude")), "owners": ws(first.get("Owner(s)")),
            "units": [{"unit_id": ws(r.get(H_UNIT)), "unit_name": ws(r.get(H_UNAME)),
                       "status": ws(r.get("Status")), "capacity_mw": ws(r.get("Capacity (MW)")),
                       "start_year": ws(r.get("Start year")),
                       "technology": ws(r.get("Turbine/Engine Technology")),
                       "chp": ws(r.get("CHP"))} for r in rows],
            "statuses": sorted({ws(r.get("Status")) for r in rows if ws(r.get("Status"))}),
            "capacity_mw": round(sum(c for c in caps if c is not None), 1),
            "captive": {h: ws(first.get(h)) for h in CAPTIVE_COLS},
            "wiki": ws(first.get("Wiki URL"))}


def looks_lng(p):
    n = p["plant_name"].lower()
    typ = p["captive"].get("Captive industry type", "").lower()
    return "lng" in n or any(t in typ for t in LIQ_TYPES)


def looks_data_center(p):
    typ = p["captive"].get("Captive industry type", "").lower()
    return "data cent" in typ


# ---------------------------------------------------------------- matching

def in_scope(row, trk, scope):
    """A sheet row is in scope when the tracker (first) or the sheet (fallback)
    puts the terminal in the scope's state or country."""
    country = (trk or {}).get("country") or SHEET_COUNTRY.get(row["sheet_country"].lower(), row["sheet_country"])
    if country.lower() != scope["country"].lower():
        return False
    if scope["kind"] == "state":
        state = (trk or {}).get("state") or row["sheet_state"]
        return state.lower() == scope["name"].lower()
    return True


def match_plants(row, trk, plants, claimed=None):
    """GEM plants that are, or may be, this terminal's captive plant, each with how
    it matched. Order: tracker ID, tracker name, name with LNG in it, close
    coordinates, then name alone or nearby coordinates. `claimed` maps a plant ID
    to the terminal whose tracker row names it: neighboring terminals (the
    Calcasieu Pass cluster sits within 3 km) must not pick up each other's
    plants by distance, and a terminal with a named plant takes no distance
    matches at all."""
    out = {}
    claimed = claimed or {}
    ids = set((trk or {}).get("captive_gem_ids") or [])
    names = {norm_key(n) for n in (trk or {}).get("captive_names") or []}
    toks = name_tokens(row["terminal"])
    tlat, tlon = (trk or {}).get("lat"), (trk or {}).get("lon")
    for pid, p in plants.items():
        how = []
        if pid in ids:
            how.append("the LNG tracker's captive plant ID column")
        if norm_key(p["plant_name"]) in names:
            how.append("the LNG tracker's captive plant column by name")
        ptoks = name_tokens(p["plant_name"])
        by_name = bool(toks) and toks <= ptoks
        d = km(tlat, tlon, p["lat"], p["lon"])
        near = d is not None and d <= NEAR_KM
        if by_name and "lng" in p["plant_name"].lower():
            how.append("the plant name")
        elif d is not None and d <= CLOSE_KM:
            how.append(f"location ({d:.1f} km from the terminal)")
        elif by_name and (near or d is None):
            # a town name alone (Freeport, Port Arthur) matches every plant in town,
            # so a name without LNG in it also has to sit by the terminal
            how.append("the plant name (the name does not say LNG; check it is the terminal's plant)")
        elif near and (p["captive"].get("Captive industry use") or "lng" in p["plant_name"].lower()):
            how.append(f"location ({d:.1f} km from the terminal; a captive plant; check it is the terminal's)")
        if not how:
            continue
        named = "ID column" in how[0] or "by name" in how[0] or how[0] == "the plant name"
        if not named and claimed.get(pid) not in (None, row["terminal_id"]):
            continue   # another terminal's tracker row names this plant
        out[pid] = {"plant": p, "matched_by": how, "distance_km": round(d, 1) if d is not None else None,
                    "named": named,
                    "confidence": "high" if named or d is not None and d <= CLOSE_KM else "medium"}
    if any(m["named"] for m in out.values()):
        out = {pid: m for pid, m in out.items() if m["named"]}
    return sorted(out.values(), key=lambda m: (m["confidence"] != "high", m["plant"]["plant_name"].lower()))


def mechanical_only(row):
    """The sheet qualified the row on compressor-drive (mechanical) turbines with
    no generating figure: GEM records generating capacity only."""
    if not row["qualifying"]:
        return False
    agg, basis = row["aggregate_mw"], row["basis"]
    if NOT_SUMMED_RE.search(agg):
        return True
    if MECH_RE.search(basis) and not re.search(r"both ways|generat|genset|electric|CHP", basis + " " + agg, re.I):
        return True
    return False


# ---------------------------------------------------------------- tasks and brief

def describe_sheet_row(row):
    parts = [f"the captive LNG sheet (tab \"{row['tab']}\", Terminal ID {row['terminal_id']})"]
    if row["qualifying"]:
        parts.append(f"lists {row['terminal']} as qualifying at the {row['threshold_mw']} MW threshold")
        if row["hardware"]:
            parts.append(f"with this equipment: {row['hardware'][:400]}")
        if row["unit_mw"]:
            parts.append(f"Per unit: {row['unit_mw'][:200]}.")
        if row["aggregate_mw"]:
            parts.append(f"In total: {row['aggregate_mw'][:200]}.")
        if row["basis"]:
            parts.append(f"Basis given: {row['basis'][:200]}.")
    else:
        parts.append(f"lists {row['terminal']} as excluded: {row['reason_excluded'][:300] or 'no reason given'}")
        if row["aggregate_mw"]:
            parts.append(f"MW figure, if any: {row['aggregate_mw'][:200]}.")
    return " ".join(parts)


def gem_line(p):
    st = ", ".join(p["statuses"]) or "no status"
    cap = p["captive"]
    cap_txt = ", ".join(f"{h.replace('Captive ', '').lower()} '{v}'" for h, v in cap.items() if v) or "captive fields blank"
    return (f"GEM has {len(p['units'])} unit(s), {st}, {p['capacity_mw']} MW in total, {cap_txt}.")


def plant_task(row, m, trk):
    p = m["plant"]
    how = m["matched_by"][0]
    refs = ", ".join(row["refs"]) or "none usable (only GEM pages were cited)"
    mech = ""
    if mechanical_only(row):
        mech = (" The sheet's figures are for turbines that drive the refrigerant compressors "
                "(mechanical drive). GEM records generating capacity only, so do not copy those "
                "figures; find the generator sets and record their electric capacity, or say that "
                "the site has no generating units at the threshold.")
    notes = ""
    if trk and trk.get("notes"):
        notes = " The LNG tracker's note says: " + " ".join(trk["notes"])[:300]
    desc = describe_sheet_row(row)
    text = (f"This plant was matched to {row['terminal']} by {how}. "
            f"{desc[:1].upper() + desc[1:]}. {gem_line(p)} Check the capacity, the status, the "
            f"start year, the technology and the captive fields against the sheet's sources and "
            f"against GEM's own sources, and report each value that differs with the page that "
            f"states it. Captive industry use is power (or both when the plant also supplies "
            f"heat) and captive industry type names LNG liquefaction; fill them when blank, with a "
            f"source.{mech} The sheet's sources: {refs}.{notes}")
    return {"plant_id": p["plant_id"], "plant_name": p["plant_name"], "unit_id": None,
            "unit_name": None, "source": "captive_lng", "section": "Captive LNG", "kind": "fix",
            "terminal": row["terminal"], "terminal_id": row["terminal_id"], "task": text,
            "fields": ["Capacity (MW)", "Status", "Start year", "Turbine/Engine Technology", "CHP",
                       "Captive industry use", "Captive industry type", "Captive non-industry use",
                       "Status Detail"],
            "checks": [6, 44], "confidence": m["confidence"]}


def render_brief(scope, data):
    name = scope["name"]
    thr = data["threshold_mw"]
    L = [f"## Captive power at LNG terminals in {name} (checklist rows 6 and 44)", ""]
    L += [f"The LNG team keeps a sheet of LNG terminals whose on-site gas-fired generation reaches "
          f"the GOGPT threshold ({thr} MW here). Read on {data['read']}: {CAPTIVE_URL}. The rows "
          f"below are the ones in {name}, joined to the LNG tracker by Terminal ID. The sheet's "
          "Qualifying rows are the only LNG captive plants at the add threshold (Baird, October "
          "2026); any other LNG captive generation you find goes to the watch list.", "",
          "Rules:", "",
          "1. Capacity is generating capacity only. Gas turbines that drive the refrigerant "
          "compressors (mechanical drive, shaft MW) are not power plant capacity. A row the "
          "sheet qualifies on mechanical drive alone needs its own generator sets found before "
          "anything is added; if there are none at the threshold, write a watch item saying so.",
          "2. Captive fields: captive industry use power (or both when heat is also supplied), "
          "captive industry type LNG liquefaction. Unit naming follows GEM's existing LNG "
          "plants (\"<Terminal name> power station\"), one unit per generator set or per "
          "combined-cycle block.",
          "3. Never cite gem.wiki; the sheet's own gem.wiki references are listed as not citable. "
          "Every other link must pass the verifier and state the value.",
          "4. Terminals under \"GEM lacks\" are new plant candidates: research them and write a "
          "full new-plant record (shards/_state.json, newplants) with \"captive_lng\": true and "
          "\"checks\": [6, 44] when the sources support generating capacity at the threshold and a "
          "concrete development step; otherwise a watch item (monitor) with the same two keys, "
          "new to the tracker, with the reason and a recheck date.",
          "5. Terminals under \"Excluded\" stay out unless you find new evidence; then a watch item.", ""]
    matched = [t for t in data["terminals"] if t["matches"]]
    lacking = [t for t in data["terminals"] if t["qualifying"] and not t["matches"]]
    excluded = [t for t in data["terminals"] if not t["qualifying"]]
    if not data["terminals"]:
        L += [f"The sheet has no row for {name}. Nothing to add from it; if an LNG terminal in "
              f"{name} has on-site generation you come across, write a watch item.", ""]
    if matched:
        L += ["### Terminals whose plant GEM already has (the plant's own brief carries the task)", ""]
        for t in matched:
            for m in t["matches"]:
                L.append(f"- {t['terminal']} ({t['terminal_id']}, {t['tracker_status'] or t['sheet_status'] or 'status not given'}) "
                         f"= GEM {m['plant_name']} ({m['plant_id']}), matched by {m['matched_by'][0]}; "
                         f"{'qualifying' if t['qualifying'] else 'excluded on the sheet'}"
                         f"{'; sheet figures are mechanical drive only' if t['mechanical_only'] else ''}.")
        L.append("")
    if lacking:
        L += ["### Qualifying terminals GEM lacks (new plant candidates)", ""]
        for t in lacking:
            L += [f"### {t['terminal']} ({t['terminal_id']})", ""]
            L.append(f"LNG tracker: {t['tracker_status'] or 'status not given'}; owner {t['owner'] or 'not given'}"
                     + (f"; parent {t['parent']}" if t["parent"] else "")
                     + (f"; operator {t['operator']}" if t["operator"] else "")
                     + (f"; at {t['location']}" if t["location"] else "")
                     + (f" ({t['lat']}, {t['lon']})" if t["lat"] is not None else "") + ".")
            L.append(f"Sheet: {t['hardware'] or 'equipment not described'}. Per unit: {t['unit_mw'] or 'not given'}. "
                     f"Total: {t['aggregate_mw'] or 'not given'}. Basis: {t['basis'] or 'not given'}.")
            if t["mechanical_only"]:
                L.append("The sheet's basis is mechanical drive only: find the generator sets before adding anything.")
            if t["captive_names"] or t["captive_gem_ids"]:
                L.append("The LNG tracker names a captive plant that is not in this scope's export: "
                         + ", ".join(t["captive_gem_ids"] + t["captive_names"]) + ". Check the tracker assignment.")
            if t["tracker_notes"]:
                L.append("LNG tracker note: " + " ".join(t["tracker_notes"])[:400])
            L.append("Sources on the sheet: " + (", ".join(t["refs"]) or "none usable"))
            if t["gem_refs"]:
                L.append("Not citable (GEM pages the sheet cited): " + ", ".join(t["gem_refs"]))
            L.append("")
    if excluded:
        L += ["### Excluded on the sheet (do not add without new evidence)", ""]
        for t in excluded:
            L.append(f"- {t['terminal']} ({t['terminal_id']}): {t['reason_excluded'] or 'no reason given'}"
                     + (f"; MW figure: {t['aggregate_mw']}" if t["aggregate_mw"] else "") + ".")
        L.append("")
    if data["gem_only"]:
        L += ["### GEM plants that look like LNG captive plants but have no sheet row", "",
              "Say in a question whether each is a captive plant at an LNG terminal the sheet "
              "missed, so the LNG team can add the row; do not change them from here.", ""]
        for p in data["gem_only"]:
            L.append(f"- {p['plant_name']} ({p['plant_id']}): {', '.join(p['statuses']) or 'no status'}, "
                     f"{p['capacity_mw']} MW, captive type '{p['captive'].get('Captive industry type') or 'blank'}'.")
        L.append("")
    if data.get("data_center") is not None:
        L += render_data_center(scope, data)
    L += ["Output: findings on tracked plants go in that plant's shard as usual. New plants and "
          "watch items from this block go in shards/_state.json under newplants and monitor, "
          "each with \"captive_lng\": true and \"checks\": [6, 44].", ""]
    return "\n".join(L)


def render_data_center(scope, data):
    name = scope["name"]
    L = [f"## Gas-powered data centers in {name} (checklist row 39, US only)", ""]
    L += ["Search for gas-fired generation built or planned to serve data centers in the state: "
          "on-site turbines or reciprocating engines behind the meter, utility plants announced "
          "for a data center campus, and existing GEM plants gaining a data center as their "
          "customer. Look in the state's air permit register and the sources in the state note, "
          "EIA-860M planned generators whose owner is a data center developer, the ISO queue, "
          "utility filings and press.", "",
          "Rules (Baird 2026-10-08):", "",
          "1. Everything from this search goes to the watch list only: shards/_state.json, "
          "monitor, never newplants or newunits. A data-center plant is added to the database "
          "in a later batch once a reviewer promotes the watch item.",
          "2. Each watch item: monitor_kind new_to_tracker when GEM has no plant (gem_plant_id "
          "blank), existing_plant with gem_plant_id when a tracked plant gains a data-center "
          "customer or an expansion for one; monitor_reason says what was found and what is "
          "still missing to add it (capacity, site, developer, a concrete step); recheck_by "
          "within six months; \"checks\": [39]; the links under refs; capacity_mw when stated.",
          "3. Emergency or backup generator sets at a data center are not captive power plants "
          "and are not watch items; only generation that supplies routine power counts.",
          "4. When the plant is later added, GEM's captive fields are captive industry use "
          "power and captive industry type data center; say so in the note so the reviewer sees it.", ""]
    dc = data["data_center"]
    if dc:
        L += [f"GEM plants in {name} already marked as data-center captive:", ""]
        for p in dc:
            L.append(f"- {p['plant_name']} ({p['plant_id']}): {', '.join(p['statuses']) or 'no status'}, "
                     f"{p['capacity_mw']} MW, owners {p['owners'] or 'blank'}.")
        L.append("")
    else:
        L += [f"GEM has no plant in {name} marked as data-center captive yet.", ""]
    return L


# ---------------------------------------------------------------- build

def make_scope(state=None, country=None):
    try:
        from build_state_brief import make_scope as _ms
        return _ms(state=state, country=country)
    except ImportError:
        if state:
            key = state.strip().lower()
            return {"kind": "state", "name": state.strip(), "slug": re.sub(r"\s+", "-", key),
                    "country": "United States", "postal": ""}
        return {"kind": "country", "name": country.strip(), "slug": re.sub(r"\s+", "-", country.strip().lower()),
                "country": country.strip(), "postal": ""}


def build(scope, sheet_rows, tracker, tracker_tab, plants, read_date):
    summaries = {pid: plant_summary(pid, rows) for pid, rows in plants.items()}
    terminals, tasks, matched_pids = [], [], set()
    thresholds = set()
    claimed = {}
    for trk in tracker.values():
        for pid in trk["captive_gem_ids"]:
            claimed.setdefault(pid, trk["terminal_id"])
        for n in trk["captive_names"]:
            for pid, p in summaries.items():
                if norm_key(p["plant_name"]) == norm_key(n):
                    claimed.setdefault(pid, trk["terminal_id"])
    for row in sheet_rows:
        trk = tracker.get(row["terminal_id"])
        if not in_scope(row, trk, scope):
            continue
        matches = match_plants(row, trk, summaries, claimed)
        t = dict(row)
        t.update({
            "in_tracker": trk is not None,
            "tracker_name": (trk or {}).get("name", ""), "tracker_status": "; ".join((trk or {}).get("statuses") or []),
            "tracker_country": (trk or {}).get("country", ""), "tracker_state": (trk or {}).get("state", ""),
            "owner": (trk or {}).get("owner", ""), "parent": (trk or {}).get("parent", ""),
            "operator": (trk or {}).get("operator", ""), "location": (trk or {}).get("location", ""),
            "lat": (trk or {}).get("lat"), "lon": (trk or {}).get("lon"),
            "captive_gem_ids": (trk or {}).get("captive_gem_ids") or [],
            "captive_names": (trk or {}).get("captive_names") or [],
            "tracker_notes": (trk or {}).get("notes") or [],
            "mechanical_only": mechanical_only(row),
            "matches": [{"plant_id": m["plant"]["plant_id"], "plant_name": m["plant"]["plant_name"],
                         "matched_by": m["matched_by"], "distance_km": m["distance_km"],
                         "confidence": m["confidence"]} for m in matches],
        })
        if row["threshold_mw"]:
            thresholds.add(row["threshold_mw"])
        for m in matches:
            matched_pids.add(m["plant"]["plant_id"])
            tasks.append(plant_task(row, m, trk))
        terminals.append(t)
    terminals.sort(key=lambda t: (not t["qualifying"], t["terminal"].lower()))
    gem_only = [p for pid, p in summaries.items() if pid not in matched_pids and looks_lng(p)]
    gem_only.sort(key=lambda p: p["plant_name"].lower())
    is_us = scope["country"].lower() == "united states"
    data_center = None
    if is_us and scope["kind"] == "state":
        data_center = sorted((p for p in summaries.values() if looks_data_center(p)),
                             key=lambda p: p["plant_name"].lower())
    eu = {"eu"}
    data = {"where": scope, "read": read_date,
            "source": {"captive_file": CAPTIVE_FILE, "captive_url": CAPTIVE_URL,
                       "lng_file": LNG_FILE, "lng_url": LNG_URL, "lng_tab": tracker_tab},
            "threshold_mw": (min(thresholds) if thresholds
                             else (20 if any(t["region"].lower() in eu for t in terminals) else 50)),
            "terminals": terminals, "gem_only": gem_only, "data_center": data_center, "tasks": tasks}
    data["brief"] = render_brief(scope, data)
    return data


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    who = ap.add_mutually_exclusive_group(required=True)
    who.add_argument("--state", action="append", default=None, help="a US state (repeatable)")
    who.add_argument("--country", action="append", default=None,
                     help="a whole country, the export's Country/Area value (repeatable)")
    ap.add_argument("--csv", default=str(gogpt_scoped_csv()), help="the GOGPT-scoped export")
    ap.add_argument("--refresh", action="store_true", help="re-download both workbooks")
    ap.add_argument("--from-dir", default=None,
                    help="a folder holding captive_lng.xlsx and lng_tracker.xlsx (offline / tests)")
    ap.add_argument("--out-dir", default=None, help="default: <repo>/work")
    a = ap.parse_args(argv)
    csv_path = Path(a.csv)
    if not csv_path.exists():
        sys.exit(f"ERROR: {csv_path} not found; run scope_filter.py first")
    cap, trk, read_date = fetch_workbooks(a.refresh, a.from_dir)
    sheet_rows = read_captive_sheet(cap)
    tracker, tracker_tab = read_lng_tracker(trk)
    missing = [r["terminal_id"] for r in sheet_rows if r["terminal_id"] not in tracker]
    if missing:
        print(f"WARNING: {len(missing)} sheet Terminal ID(s) not in the LNG tracker download: "
              f"{', '.join(missing[:8])}{' ...' if len(missing) > 8 else ''}; their country and state "
              "come from the sheet alone")
    out_dir = Path(a.out_dir) if a.out_dir else work_dir()
    out_dir.mkdir(parents=True, exist_ok=True)
    scopes = ([make_scope(state=ws(s)) for s in a.state] if a.state
              else [make_scope(country=ws(c)) for c in a.country])
    for scope in scopes:
        plants = load_export(csv_path, scope)
        data = build(scope, sheet_rows, tracker, tracker_tab, plants, read_date)
        tag = f"us-{scope['slug']}" if scope["kind"] == "state" else scope["slug"]
        jp, mp = out_dir / f"captive_{tag}.json", out_dir / f"captive_{tag}.md"
        jp.write_text(json.dumps(data, indent=2, ensure_ascii=False), encoding="utf-8")
        mp.write_text(data["brief"] + "\n", encoding="utf-8")
        terms = data["terminals"]
        q = [t for t in terms if t["qualifying"]]
        lacking = [t for t in q if not t["matches"]]
        print(f"{scope['name']}: {len(terms)} sheet row(s) in scope ({len(q)} qualifying at "
              f"{data['threshold_mw']} MW, {len(terms) - len(q)} excluded); {len(data['tasks'])} plant "
              f"task(s) on {len({t['plant_id'] for t in data['tasks']})} GEM plant(s); {len(lacking)} qualifying "
              f"terminal(s) GEM lacks; {len(data['gem_only'])} GEM LNG-looking plant(s) with no sheet row"
              + (f"; {len(data['data_center'])} data-center plant(s) already in GEM" if data["data_center"] is not None else "")
              + f". Wrote {jp} and {mp}")
        for t in terms:
            if t["matches"]:
                names = "; ".join(f"{m['plant_name']} ({m['plant_id']}, {m['matched_by'][0]})" for m in t["matches"])
            else:
                names = "no GEM plant matched"
            flag = " [mechanical drive only]" if t["mechanical_only"] else ""
            print(f"  {'Q' if t['qualifying'] else 'X'} {t['terminal']} ({t['terminal_id']}){flag}: {names}")


if __name__ == "__main__":
    main()
